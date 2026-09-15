#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
split_pickup_by_yard.py
------------------------------------------------------------------------
อ่านไฟล์ Excel (Pickup List / BKG) ทุกไฟล์ในโฟลเดอร์ แล้ว "แยกข้อมูลตามลานลากตู้"
(คอลัมน์ Pickup / Pickup Name) ออกมาเป็นไฟล์ Excel แยกกันคนละไฟล์ต่อ 1 ลาน
+ ไฟล์ Summary รวมทุกลาน

แนวคิด: ดึงข้อมูลด้วยไลบรารีล้วน ๆ ไม่ใช้ AI
    - อ่าน .xls  ด้วย pandas + xlrd
    - อ่าน .xlsx ด้วย pandas + openpyxl
    - ถ้ามีไฟล์ .pdf ในโฟลเดอร์ จะใช้ pdfplumber ดึง text ออกมาเป็นไฟล์ .txt ให้ด้วย

ผลลัพธ์ต่อ 1 ไฟล์ต้นทาง (<stem> = ชื่อไฟล์ต้นทางไม่รวมนามสกุล):
    <stem> - Summary.xlsx          -> ชีต "Summary"  : 1 แถวต่อ 1 ลาน + แถว TOTAL
                                       + ตาราง Group Summary (BKK / LCH, LCH55 นับรวม BKK)
    <stem> - <ชื่อลาน>.xlsx         -> ชีต "Pickup List" : booking ทั้งหมดของลานนั้น
                                       แถวแรก = ชื่อเรื่อง, แถว 2 = หัวตาราง, ท้าย = TOTAL (n bookings)

การจัดรูปแบบ (ให้เหมือนไฟล์ตัวอย่างของผู้ใช้):
    - ฟอนต์ Calibri 11 ทั้งไฟล์ (ชื่อเรื่อง 13 ตัวหนา), ความสูงแถว 13 ทุกแถว
    - หัวตารางพื้นน้ำเงินเข้ม (1F4E78) ตัวอักษรขาว, freeze ที่ A3
    - แถว TOTAL พื้นฟ้าอ่อน (D9E1F2) ตัวหนา + สูตร =SUM(...)
    - ไฮไลท์ทั้งแถว: เหลืองถ้า TPSZ มี 20RF / R40H, เขียวถ้ามี 20OT / 20FR / 40OT / 40FR
    - TRAFFIC ORDER ที่มีคำว่า precool / pre-cool -> ตัวอักษรแดงตัวหนา
    - เติม TRAFFIC ORDER / ORG CUST ที่ว่าง, ตัดคอลัมน์ DOC CUST / Pickup Name
    - COMMON REMARK: ตัดออก ยกเว้นลาน BKK01 / BKK02 / BKK04 / LCH55
    - ETD / TRAN DT เก็บรูปแบบตัวเลขดิบ (เช่น 202609190800, 20260911)

วิธีใช้:
    python split_pickup_by_yard.py                       # อ่านโฟลเดอร์ที่สคริปต์อยู่
    python split_pickup_by_yard.py "D:\\path\\to\\folder"  # ระบุโฟลเดอร์
    python split_pickup_by_yard.py "D:\\folder" -o "D:\\out"

ต้องมีไลบรารี: pandas, xlrd (อ่าน .xls), openpyxl (เขียน .xlsx)  [pdfplumber = ถ้ามี]
    pip install pandas xlrd openpyxl pdfplumber
------------------------------------------------------------------------
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# ตั้งค่า
# ----------------------------------------------------------------------

YARD_CODE_COL = "Pickup"        # รหัสลาน เช่น BKK27, LCHY5
YARD_NAME_COL = "Pickup Name"   # ชื่อลาน เช่น B.C. DEPOT CO.,LTD.

# รวมลานหลายรหัสให้ออกเป็นไฟล์เดียว + แถวเดียวในชีต Summary
#   merge_id: {"codes": [...เรียงลำดับที่อยากให้แสดง...], "file_label": "ชื่อสั้นสำหรับตั้งชื่อไฟล์"}
#   - ชื่อในหัวเรื่อง / ชีต Summary  = ชื่อลานจริงของแต่ละรหัส ต่อกันด้วย " + "
#   - เซลล์ Pickup ในชีต Summary    = รหัสต่อกันด้วย " + "  (เช่น "BKK01 + BKK04")
#   - ชื่อไฟล์                       = "<file_label> (<codes คั่นด้วย _>)"
YARD_MERGE = {
    "PAT_BKK": {
        "codes": ["BKK01", "BKK04"],
        "file_label": "PAT TERMINAL 1-2 (PORT AUTHORITY OF THAILAND)",
    },
}

# ตัดช่องว่าง/ขึ้นบรรทัดหน้า-หลัง เฉพาะคอลัมน์ข้อความยาว (คอลัมน์อื่น เช่น TPSZ คงค่าดิบ)
STRIP_COLS = {"ORG CUST", "COMMODITY", "TRAFFIC ORDER", "COMMON REMARK"}

# ถ้าช่องเหล่านี้ "ว่างจริง" (None / NaN / ช่องว่างล้วน) ให้เติมข้อความนี้อัตโนมัติ
FILL_WHEN_BLANK = {
    "TRAFFIC ORDER": "GOOD AND CLEAN CONTAINER",
}

# COMMON REMARK: ตัดออกจากไฟล์ลาน ยกเว้นลาน (รหัส Pickup เดิม) เหล่านี้
COMMON_REMARK_KEEP = {"BKK01", "BKK02", "BKK04", "LCH55"}

# ในชีต Summary: ลานเหล่านี้ให้นับรวมอยู่กลุ่ม BKK
GROUP_BKK_EXTRA = {"LCH55"}

# ----------------------------------------------------------------------
# คอลัมน์ในไฟล์ลาน (ชีต "Pickup List")  — ไม่มี Pickup / Pickup Name / DOC CUST
# ----------------------------------------------------------------------
YARD_COLUMNS = [
    "BK No", "VSL", "VOY", "ETD", "POR", "LOD", "DIS", "TPSZ", "TRAN DT",
    "ORG CUST", "COMMODITY", "TRAFFIC ORDER", "COMMON REMARK",
    "GP22", "GP42", "GP45", "RE22", "RE45", "UT22", "UT42", "PC22", "PC42",
    "Pickup Qty", "Return",
]
COL_DISPLAY = {"Pickup Qty": "Pickup"}   # ชื่อที่แสดงในหัวตาราง

# ความกว้างคอลัมน์ (ตามชื่อที่แสดง); คอลัมน์ที่ไม่ระบุ = ค่าเริ่มต้นของ Excel
COL_WIDTH = {
    "BK No": 20, "VSL": 8, "ETD": 14, "POR": 6, "LOD": 8, "TPSZ": 12,
    "ORG CUST": 32, "COMMODITY": 20, "TRAFFIC ORDER": 16, "COMMON REMARK": 40,
    "GP22": 7, "Pickup": 9,
}

QTY_COLS = ["GP22", "GP42", "GP45", "RE22", "RE45", "UT22", "UT42", "PC22", "PC42"]

# ----------------------------------------------------------------------
# สไตล์
# ----------------------------------------------------------------------
FONT_NAME = "Calibri"
FONT_SIZE = 11
TITLE_SIZE = 13
GROUP_TITLE_SIZE = 12
ROW_HEIGHT = 13

HEADER_FILL = "FF1F4E78"   # หัวตาราง (ตัวอักษรขาว)
TOTAL_FILL = "FFD9E1F2"    # แถวรวม
YELLOW_FILL = "FFFFFF00"
GREEN_FILL = "FF92D050"
WHITE = "FFFFFFFF"
RED = "FFFF0000"

TPSZ_YELLOW = ("20RF", "R40H")                     # ตู้ห้องเย็น
TPSZ_GREEN = ("20OT", "20FR", "40OT", "40FR")      # ตู้ open-top / flat-rack
PRECOOL_RE = re.compile(r"pre[-\s]?cool", re.IGNORECASE)

_BASE_FONT = Font(name=FONT_NAME, size=FONT_SIZE)
_BOLD_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
_HEAD_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=WHITE)
_TITLE_FONT = Font(name=FONT_NAME, size=TITLE_SIZE, bold=True)
_GTITLE_FONT = Font(name=FONT_NAME, size=GROUP_TITLE_SIZE, bold=True)
_PRECOOL_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=RED)
_HEAD_PAT = PatternFill("solid", fgColor=HEADER_FILL)
_TOTAL_PAT = PatternFill("solid", fgColor=TOTAL_FILL)
_YELLOW_PAT = PatternFill("solid", fgColor=YELLOW_FILL)
_GREEN_PAT = PatternFill("solid", fgColor=GREEN_FILL)


# ----------------------------------------------------------------------
# ตัวช่วยทั่วไป
# ----------------------------------------------------------------------

def _norm(x) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    return str(x).strip()


def _is_blank(x) -> bool:
    if x is None:
        return True
    if isinstance(x, float) and pd.isna(x):
        return True
    return isinstance(x, str) and x.strip() == ""


def _cell(v):
    """แปลงค่าให้ openpyxl เขียนได้ (numpy scalar -> python, NaN/NaT -> None)"""
    if isinstance(v, str):
        return v
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(v, "item"):
        try:
            return v.item()
        except Exception:
            pass
    return v


def sanitize_filename(text: str) -> str:
    """ตัดอักขระที่ตั้งชื่อไฟล์ Windows ไม่ได้ออก (comma -> _ ให้เหมือนไฟล์ตัวอย่าง)"""
    text = re.sub(r'[\\/:*?"<>|,]', "_", str(text))
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s{2,}", " ", text).strip()
    return text or "NA"


def make_unique_columns(names) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for n in names:
        n = "Unnamed" if (n is None or (isinstance(n, float) and pd.isna(n))) else str(n).strip()
        if n in seen:
            seen[n] += 1
            out.append(f"{n}.{seen[n]}")
        else:
            seen[n] = 0
            out.append(n)
    return out


def digits_only(value) -> str | None:
    """'202609230130.0' -> '202609230130' ; '20260911' -> '20260911'"""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    s = str(value).strip()
    if not s:
        return None
    s = s.split(".")[0]
    s = re.sub(r"\D", "", s)
    return s or None


# ----------------------------------------------------------------------
# อ่าน Excel
# ----------------------------------------------------------------------

def find_header_row(raw: pd.DataFrame, key: str = "BK No", scan_rows: int = 20) -> int:
    limit = min(scan_rows, len(raw))
    for i in range(limit):
        for cell in raw.iloc[i].tolist():
            if isinstance(cell, str) and cell.strip().lower() == key.lower():
                return i
    return 0


def read_pickup_table(path: Path) -> pd.DataFrame:
    raw = pd.read_excel(path, header=None, dtype=object)
    hdr = find_header_row(raw)
    header = make_unique_columns(raw.iloc[hdr].tolist())
    df = raw.iloc[hdr + 1:].copy()
    df.columns = header
    df = df.reset_index(drop=True)

    if "Pickup.1" in df.columns:
        df = df.rename(columns={"Pickup.1": "Pickup Qty"})

    if "BK No" in df.columns:
        bk = df["BK No"].map(lambda x: "" if x is None or (isinstance(x, float) and pd.isna(x)) else str(x).strip())
        df = df[bk.ne("") & bk.str.lower().ne("nan") & bk.str.lower().ne("none")].copy()

    df = df.dropna(axis=1, how="all")
    return df.reset_index(drop=True)


def clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    """เติม / ตัด / จัดชนิดข้อมูล (ไม่แปลง ETD/TRAN DT เป็นรูปแบบวันที่ — คงตัวเลขดิบ)"""
    df = df.copy()

    # ตัดช่องว่างหน้า-หลัง เฉพาะคอลัมน์ข้อความยาว (คงค่าดิบของคอลัมน์อื่น เช่น TPSZ = "40Hx1 ")
    for col in STRIP_COLS:
        if col in df.columns:
            df[col] = df[col].map(lambda x: x.strip() if isinstance(x, str) else x)

    # เติมช่องที่ "ว่างจริง" เท่านั้น
    for col, value in FILL_WHEN_BLANK.items():
        if col in df.columns:
            df[col] = df[col].map(lambda x: value if _is_blank(x) else x)

    # ORG CUST ว่าง -> เติมจาก DOC CUST แล้วตัด DOC CUST ทิ้ง
    if "ORG CUST" in df.columns and "DOC CUST" in df.columns:
        blank = df["ORG CUST"].map(_is_blank)
        df.loc[blank, "ORG CUST"] = df.loc[blank, "DOC CUST"]
    if "DOC CUST" in df.columns:
        df = df.drop(columns=["DOC CUST"])

    # ETD / TRAN DT / Pickup Qty / Return -> ตัวเลขดิบเป็น string
    for col in ("ETD", "TRAN DT", "Pickup Qty", "Return"):
        if col in df.columns:
            df[col] = df[col].map(lambda v: digits_only(v) or _norm(v))

    # คอลัมน์จำนวนตู้ -> จำนวนเต็ม
    for col in QTY_COLS:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0).astype(int)

    return df


# ----------------------------------------------------------------------
# การจัดกลุ่มลาน
# ----------------------------------------------------------------------

_CODE2MERGE = {c.upper(): mid for mid, cfg in YARD_MERGE.items() for c in cfg["codes"]}


def yard_key(code: str) -> str:
    u = (code or "").upper()
    return _CODE2MERGE.get(u, u or "NA")


def group_of(disp_code: str) -> str:
    first = disp_code.split(" + ")[0].strip().upper()
    if first in GROUP_BKK_EXTRA or first.startswith("BKK"):
        return "BKK"
    if first.startswith("LCH"):
        return "LCH"
    return "OTHER"


# ----------------------------------------------------------------------
# เขียน Excel
# ----------------------------------------------------------------------

def _set_row_heights(ws, last_row: int) -> None:
    for r in range(1, last_row + 1):
        ws.row_dimensions[r].height = ROW_HEIGHT


def write_yard_file(out_path: Path, g: pd.DataFrame, title_name: str, keep_common_remark: bool) -> None:
    if "BK No" in g.columns:
        g = g.sort_values("BK No", kind="stable")
    cols = [c for c in YARD_COLUMNS if c in g.columns]
    if not keep_common_remark and "COMMON REMARK" in cols:
        cols.remove("COMMON REMARK")
    disp = [COL_DISPLAY.get(c, c) for c in cols]
    ncol = len(cols)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pickup List"

    # แถว 1 : ชื่อเรื่อง
    ws.cell(1, 1, f"Pickup List - {title_name}").font = _TITLE_FONT

    # แถว 2 : หัวตาราง
    for j, name in enumerate(disp, start=1):
        c = ws.cell(2, j, name)
        c.font = _HEAD_FONT
        c.fill = _HEAD_PAT

    # แถว 3+ : ข้อมูล
    tpsz_i = cols.index("TPSZ") if "TPSZ" in cols else None
    traffic_i = cols.index("TRAFFIC ORDER") if "TRAFFIC ORDER" in cols else None
    qty_pos = [cols.index(q) for q in QTY_COLS if q in cols]
    first_qty = (min(qty_pos) + 1) if qty_pos else ncol + 1

    n = len(g)
    for i, (_, rec) in enumerate(g.iterrows()):
        r = 3 + i
        tp = _norm(rec.get("TPSZ")).upper().replace(" ", "")
        fill = None
        if any(k in tp for k in TPSZ_YELLOW):
            fill = _YELLOW_PAT
        elif any(k in tp for k in TPSZ_GREEN):
            fill = _GREEN_PAT
        precool = traffic_i is not None and bool(PRECOOL_RE.search(_norm(rec.get("TRAFFIC ORDER"))))

        for j, col in enumerate(cols, start=1):
            cell = ws.cell(r, j, _cell(rec[col]))
            cell.font = _BASE_FONT
            if fill is not None:
                cell.fill = fill
        if precool:
            ws.cell(r, traffic_i + 1).font = _PRECOOL_FONT

    last_data = 2 + n
    total_r = last_data + 1

    # แถว TOTAL — จัดสไตล์ก่อน แล้วค่อย merge
    for j in range(1, ncol + 1):
        cell = ws.cell(total_r, j)
        cell.font = _BASE_FONT
        cell.fill = _TOTAL_PAT
    a1 = ws.cell(total_r, 1, f"TOTAL ({n} bookings)")
    a1.font = _BOLD_FONT
    for q in QTY_COLS:
        if q in cols:
            j = cols.index(q) + 1
            L = get_column_letter(j)
            c = ws.cell(total_r, j, f"=SUM({L}3:{L}{last_data})" if n else 0)
            c.font = _BOLD_FONT
            c.fill = _TOTAL_PAT
    if first_qty > 2:
        ws.merge_cells(start_row=total_r, start_column=1, end_row=total_r, end_column=first_qty - 1)

    # ความกว้างคอลัมน์
    for j, name in enumerate(disp, start=1):
        if name in COL_WIDTH:
            ws.column_dimensions[get_column_letter(j)].width = COL_WIDTH[name]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.freeze_panes = "A3"
    _set_row_heights(ws, total_r)
    wb.save(out_path)


def write_summary_file(out_path: Path, stem: str, yard_rows: list[dict]) -> None:
    head = ["Group", "Pickup", "Pickup Name", "Number of Records"] + QTY_COLS + ["Total Containers"]
    ncol = len(head)  # 14
    rank = {"BKK": 0, "LCH": 1, "OTHER": 2}
    yard_rows = sorted(yard_rows, key=lambda x: (rank.get(x["group"], 9), x["disp_code"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws.cell(1, 1, f"Pickup Name Summary - {stem}").font = _TITLE_FONT
    for j, h in enumerate(head, start=1):
        c = ws.cell(2, j, h)
        c.font = _HEAD_FONT
        c.fill = _HEAD_PAT

    r0 = 3
    for i, y in enumerate(yard_rows):
        r = r0 + i
        vals = ([y["group"], y["disp_code"], y["disp_name"], y["n_records"]]
                + [y["qty"][q] for q in QTY_COLS] + [y["total"]])
        for j, v in enumerate(vals, start=1):
            ws.cell(r, j, _cell(v)).font = _BASE_FONT
    last = r0 + len(yard_rows) - 1
    total_r = last + 1

    for j in range(1, ncol + 1):
        cell = ws.cell(total_r, j)
        cell.font = _BOLD_FONT
        cell.fill = _TOTAL_PAT
    ws.cell(total_r, 1, "TOTAL")
    for j in range(4, ncol + 1):
        L = get_column_letter(j)
        ws.cell(total_r, j, f"=SUM({L}{r0}:{L}{last})")
    ws.merge_cells(start_row=total_r, start_column=1, end_row=total_r, end_column=3)

    # ---- Group Summary ----
    gt = total_r + 2
    ws.cell(gt, 3, "Group Summary (LCH55 counted under BKK)").font = _GTITLE_FONT
    gh = gt + 1
    ghead = ["Group", "Number of Records"] + QTY_COLS + ["Total Containers"]
    for k, h in enumerate(ghead):
        c = ws.cell(gh, 3 + k, h)
        c.font = _HEAD_FONT
        c.fill = _HEAD_PAT

    groups = [g for g in ("BKK", "LCH", "OTHER") if any(y["group"] == g for y in yard_rows)] or ["BKK", "LCH"]
    grows = []
    for gi, gname in enumerate(groups):
        gr = gh + 1 + gi
        ws.cell(gr, 3, gname).font = _BOLD_FONT
        for j in range(4, ncol + 1):
            L = get_column_letter(j)
            ws.cell(gr, j, f'=SUMIF(A{r0}:A{last},"{gname}",{L}{r0}:{L}{last})').font = _BASE_FONT
        grows.append(gr)
    gtot = grows[-1] + 1
    for j in range(3, ncol + 1):
        cell = ws.cell(gtot, j)
        cell.font = _BOLD_FONT
        cell.fill = _TOTAL_PAT
    ws.cell(gtot, 3, "TOTAL")
    for j in range(4, ncol + 1):
        L = get_column_letter(j)
        ws.cell(gtot, j, f"=SUM({L}{grows[0]}:{L}{grows[-1]})")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.merge_cells(start_row=gt, start_column=3, end_row=gt, end_column=ncol)

    for col, w in {"A": 10, "B": 22, "C": 40, "D": 9}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    _set_row_heights(ws, gtot)
    wb.save(out_path)


# ----------------------------------------------------------------------
# pdfplumber (ถ้ามีไฟล์ PDF ในโฟลเดอร์)
# ----------------------------------------------------------------------

def dump_pdf_text(folder: Path, out_dir: Path) -> None:
    pdfs = sorted(folder.glob("*.pdf"))
    if not pdfs:
        return
    try:
        import pdfplumber
    except ImportError:
        print(f"  ! พบไฟล์ PDF {len(pdfs)} ไฟล์ แต่ยังไม่ได้ติดตั้ง pdfplumber (ข้าม) -> pip install pdfplumber")
        return
    for pdf_path in pdfs:
        txt_path = out_dir / f"{pdf_path.stem}__text.txt"
        parts: list[str] = []
        with pdfplumber.open(pdf_path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                parts.append(f"\n===== หน้า {i} / {pdf_path.name} =====\n")
                parts.append(page.extract_text() or "")
        txt_path.write_text("\n".join(parts), encoding="utf-8")
        print(f"  - ดึงข้อความจาก PDF: {pdf_path.name} -> {txt_path.name}")


# ----------------------------------------------------------------------
# main
# ----------------------------------------------------------------------

def process_excel_file(path: Path, out_dir: Path) -> None:
    print(f"\n[ไฟล์] {path.name}")
    df = read_pickup_table(path)
    if df.empty:
        print("  ! ไม่พบข้อมูลในไฟล์นี้ (ข้าม)")
        return
    if YARD_CODE_COL not in df.columns and YARD_NAME_COL not in df.columns:
        print(f"  ! ไม่พบคอลัมน์ '{YARD_CODE_COL}' หรือ '{YARD_NAME_COL}' (ข้าม)")
        return

    clean = clean_frame(df)
    print(f"  อ่านได้ {len(clean)} แถว, {len(clean.columns)} คอลัมน์")
    stem = path.stem

    codes = (clean[YARD_CODE_COL] if YARD_CODE_COL in clean else pd.Series([""] * len(clean))).map(_norm)
    names = (clean[YARD_NAME_COL] if YARD_NAME_COL in clean else pd.Series([""] * len(clean))).map(_norm)
    clean = clean.assign(_key=codes.map(yard_key).values, _code=codes.values, _name=names.values)

    # รหัสลาน -> ชื่อลานตัวแทน (ที่พบบ่อยสุด)
    code_name: dict[str, str] = {}
    for cd, gg in clean.groupby("_code"):
        if not cd:
            continue
        m = gg.loc[gg["_name"] != "", "_name"]
        code_name[cd] = m.mode().iat[0] if not m.mode().empty else ""

    yard_rows: list[dict] = []
    for key, g in clean.groupby("_key"):
        present = [c for c in dict.fromkeys(g["_code"]) if c]
        mcfg = YARD_MERGE.get(key)
        if mcfg:
            present_up = {p.upper() for p in present}
            ordered = [c for c in mcfg["codes"] if c.upper() in present_up] or list(mcfg["codes"])
            disp_code = " + ".join(ordered)
            disp_name = " + ".join(code_name.get(c, c) for c in ordered)
            file_label = f'{mcfg["file_label"]} ({"_".join(ordered)})'
            keep_cr = bool({c.upper() for c in ordered} & COMMON_REMARK_KEEP)
            title_name = disp_name
        elif not present:
            disp_code, disp_name = "NA", ""
            file_label, keep_cr, title_name = "NA", False, "(ไม่ระบุลาน)"
        else:
            cd = present[0]
            disp_code = cd
            disp_name = code_name.get(cd, cd)
            file_label = sanitize_filename(disp_name) if disp_name else cd
            keep_cr = cd.upper() in COMMON_REMARK_KEEP
            title_name = disp_name or cd

        qsum = {q: int(g[q].sum()) if q in g else 0 for q in QTY_COLS}
        yard_rows.append({
            "group": group_of(disp_code),
            "disp_code": disp_code,
            "disp_name": disp_name,
            "n_records": len(g),
            "qty": qsum,
            "total": sum(qsum.values()),
        })

        yard_path = out_dir / f"{stem} - {file_label}.xlsx"
        write_yard_file(yard_path, g, title_name, keep_cr)
        print(f"  + {yard_path.name}  ({len(g)} bookings)")

    summary_path = out_dir / f"{stem} - Summary.xlsx"
    write_summary_file(summary_path, stem, yard_rows)
    print(f"  + {summary_path.name}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="แยกข้อมูล Pickup List ตามลานลากตู้ ออกเป็นไฟล์ Excel คนละไฟล์ + Summary"
    )
    parser.add_argument("folder", nargs="?", default=None,
                        help="โฟลเดอร์ที่มีไฟล์ Excel (ค่าเริ่มต้น = โฟลเดอร์ที่สคริปต์นี้อยู่)")
    parser.add_argument("-o", "--output", default=None,
                        help="โฟลเดอร์ผลลัพธ์ (ค่าเริ่มต้น = <folder>/output)")
    args = parser.parse_args(argv)

    folder = Path(args.folder).expanduser().resolve() if args.folder else Path(__file__).parent.resolve()
    if not folder.is_dir():
        print(f"ไม่พบโฟลเดอร์: {folder}")
        return 1

    out_dir = Path(args.output).expanduser().resolve() if args.output else (folder / "output")
    out_dir.mkdir(parents=True, exist_ok=True)

    excel_files = []
    for pattern in ("*.xls", "*.xlsx", "*.xlsm"):
        for p in folder.glob(pattern):
            if p.name.startswith("~$"):
                continue
            if out_dir in p.parents or p.parent == out_dir:
                continue
            excel_files.append(p)
    excel_files = sorted(set(excel_files))

    print(f"โฟลเดอร์ต้นทาง : {folder}")
    print(f"โฟลเดอร์ผลลัพธ์: {out_dir}")
    print(f"พบไฟล์ Excel   : {len(excel_files)} ไฟล์")

    if not excel_files:
        print("ไม่พบไฟล์ Excel ในโฟลเดอร์นี้")
    for path in excel_files:
        try:
            process_excel_file(path, out_dir)
        except Exception as exc:  # noqa: BLE001
            print(f"  ! ผิดพลาดกับไฟล์ {path.name}: {exc}")

    dump_pdf_text(folder, out_dir)
    print("\nเสร็จแล้ว")
    return 0


if __name__ == "__main__":
    sys.exit(main())
